# backend/main.py
import os
import json
import fitz  # PyMuPDF
import logging
import io
from docx import Document
from pptx import Presentation
from flask import Flask, request, jsonify
from flask_cors import CORS
from waitress import serve
from dotenv import load_dotenv
from datetime import datetime, timedelta

# --- Imports for DB and Auth ---
from flask_sqlalchemy import SQLAlchemy
from flask_bcrypt import Bcrypt
from flask_jwt_extended import create_access_token, get_jwt_identity, JWTManager, verify_jwt_in_request


# --- Google GenAI Library Imports ---
from google import genai
from google.genai import types
from google.api_core import exceptions as google_exceptions

# --- App Initialization & Config ---
load_dotenv()
app = Flask(__name__)

# Configure logging
logging.basicConfig(level=logging.INFO)
app.logger.setLevel(logging.INFO)

# In production, use a secure, randomly generated secret key stored as an env variable
app.config["JWT_SECRET_KEY"] = os.getenv("JWT_SECRET_KEY")
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///site.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["JWT_ACCESS_TOKEN_EXPIRES"] = timedelta(hours=1)
app.config["JWT_CSRF_PROTECTION"] = False
app.config["JWT_CSRF_IN_COOKIES"] = False

CORS(app, supports_credentials=True)

# --- Initialize Extensions ---
db = SQLAlchemy(app)
bcrypt = Bcrypt(app)
jwt = JWTManager(app)

# --- Database Models ---
class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(60), nullable=False)
    subscription_tier = db.Column(db.String(20), nullable=False, default='free')

class UsageLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    timestamp = db.Column(db.DateTime, nullable=False, default=datetime.now)

# --- Tier Limits Configuration ---
TIER_LIMITS = {
    'free': 3,
    'pro': 50
}

# --- Master System Prompt for AI ---
SYSTEM_PROMPT = """
You are an expert curriculum designer and a helpful AI assistant for teachers. Your task is to generate high-quality, relevant quiz questions based on the provided text document. The output MUST be a valid JSON array of question objects. Do not include any explanatory text, notes, or markdown formatting like ```json ... ``` before or after the JSON array. The response must start with '[' and end with ']'. Each object in the array must conform to one of the following structures based on the 'type' key. For "single" choice questions: {"id": "placeholder_id", "type": "single", "Question": "The question text?", "Options": ["Option A", "Correct Option B", "Option C", "Option D"], "answer": "Correct Option B", "Rationale": "A brief explanation of why the answer is correct.", "hint": "An optional hint for the student."}. For "multi-select" questions: {"id": "placeholder_id", "type": "multi-select", "Question": "The question text, asking for multiple answers?", "Options": ["Correct Option A", "Incorrect Option B", "Correct Option C", "Incorrect Option D"], "answer": ["Correct Option A", "Correct Option C"], "Rationale": "A brief explanation of why the selected answers are correct.", "hint": "An optional hint for the student."}. For "fill-in" questions: {"id": "placeholder_id", "type": "fill-in", "Question": "The capital of France is ____.", "Options": [], "answer": "Paris", "Rationale": "Paris is the capital city of France.", "hint": "It's a famous European city known for art."}. For "ordering" questions: {"id": "placeholder_id", "type": "ordering", "Question": "Arrange these planets in order from the sun.", "Options": ["Mars", "Venus", "Earth", "Mercury"], "answer": ["Mercury", "Venus", "Earth", "Mars"], "Rationale": "This is the correct order of the first four planets from the sun.", "hint": "A hot planet is first."}
"""

# --- API Endpoints ---
@app.route('/api/register', methods=['POST'])
def register():
    app.logger.info("/api/register endpoint hit.")
    data = request.get_json()
    email = data.get('email')
    password = data.get('password')

    if not email or not password:
        app.logger.warning("Registration attempt with missing email or password.")
        return jsonify({"message": "Email and password are required."}), 400

    if User.query.filter_by(email=email).first():
        app.logger.warning(f"Registration attempt for existing email: {email}")
        return jsonify({"message": "Email already exists"}), 409
    
    app.logger.info(f"Registering new user: {email}")
    hashed_password = bcrypt.generate_password_hash(password).decode('utf-8')
    user = User(email=email, password_hash=hashed_password)
    db.session.add(user)
    db.session.commit()
    app.logger.info(f"User {email} registered successfully.")
    return jsonify({"message": "User registered successfully"}), 201

@app.route('/api/login', methods=['POST'])
def login():
    app.logger.info("/api/login endpoint hit.")
    data = request.get_json()
    email = data.get('email')
    password = data.get('password')

    if not email or not password:
        app.logger.warning("Login attempt with missing email or password.")
        return jsonify({"message": "Email and password are required."}), 400

    user = User.query.filter_by(email=email).first()
    if user and bcrypt.check_password_hash(user.password_hash, password):
        app.logger.info(f"Successful login for user: {user.email}")
        access_token = create_access_token(identity=str(user.id))
        return jsonify(access_token=access_token)
    
    app.logger.warning(f"Failed login attempt for email: {email}")
    return jsonify({"message": "Invalid credentials"}), 401


@app.route('/api/generate-questions', methods=['POST'])
def generate_questions():
    app.logger.info("/api/generate-questions endpoint hit.")
    
    # --- Step 1: Manual Token Verification ---
    try:
        verify_jwt_in_request()
    except Exception as e:
        app.logger.warning(f"JWT verification failed: {str(e)}")
        return jsonify(message="Missing or invalid token"), 401

    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    if not user:
        return jsonify({"message": "User not found"}), 404
    app.logger.info(f"Request token VERIFIED for user: {user.email} (ID: {user.id})")

    # --- Step 2: Usage Check ---
    is_dev_user = os.getenv("DEV_MODE_USER_EMAIL") == user.email
    if not is_dev_user:
        start_of_month = datetime.utcnow().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        usage_count = UsageLog.query.filter(UsageLog.user_id == user.id, UsageLog.timestamp >= start_of_month).count()
        limit = TIER_LIMITS.get(user.subscription_tier, 0)
        if usage_count >= limit:
            app.logger.warning(f"User {user.email} has reached their generation limit.")
            return jsonify({"success": False, "message": f"You have used {usage_count}/{limit} of your monthly generations. Please upgrade for more."}), 429

    # --- Step 3: Load API Key ---
    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        app.logger.error("CRITICAL: GOOGLE_API_KEY not configured on server.")
        return jsonify({"message": "Server configuration error: API key not found."}), 500

    # --- Step 4: File and Form Processing ---
    try:
        if 'document' not in request.files: return jsonify({"success": False, "message": "No document file provided."}), 400
        file = request.files['document']
        if file.filename == '': return jsonify({"success": False, "message": "No selected file."}), 400

        subject = request.form.get('subject', 'General')
        grade = request.form.get('grade', 'Unspecified')
        notes = request.form.get('notes', 'None')
        num_questions = request.form.get('num_questions', '5')
        question_types = request.form.get('question_types', 'single, multi-select')

        file_bytes = file.read()
        filename_lower = file.filename.lower()
        
        # This list will be populated differently depending on file type
        contents_for_api = []

        if filename_lower.endswith('.pptx'):
            app.logger.info(f"Processing PPTX file: {file.filename} using hybrid image/text model.")
            # Use io.BytesIO to treat the byte stream like a file, avoiding saving to disk
            pptx_stream_for_text = io.BytesIO(file_bytes)
            pptx_stream_for_images = io.BytesIO(file_bytes)

            prs = Presentation(pptx_stream_for_text)
            doc_for_images = fitz.open(stream=pptx_stream_for_images, filetype="pptx")
            
            # For multimodal requests, the system prompt and instructions go first in the list
            initial_prompt = f"{SYSTEM_PROMPT}\n\nPlease generate {num_questions} questions based on the content of the following presentation slides. Each slide is provided as both an image and its extracted text. Use both to understand the full context.\n- Subject: {subject}\n- Grade Level: {grade}\n- Desired Question Types: {question_types}\n- Additional Teacher Notes: {notes}\n---"
            contents_for_api.append(initial_prompt)

            # Process each slide for its image and text
            for i, slide in enumerate(prs.slides):
                # Part 1: The Image
                page_for_image = doc_for_images.load_page(i)
                pix = page_for_image.get_pixmap(dpi=150)
                img_bytes = pix.tobytes("png")
                image_part = types.Part.from_bytes(data=img_bytes, mime_type='image/png')
                contents_for_api.append(image_part)

                # Part 2: The Extracted Text from the same slide
                slide_text_parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text:
                        slide_text_parts.append(shape.text_frame.text.strip())
                
                # Also extract speaker notes for additional context
                notes_text = ""
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
                
                # Combine text and notes into a single text part for the AI
                text_content_for_slide = f"\n--- Extracted text for the slide above (Slide {i+1}) ---\n" + "\n".join(slide_text_parts)
                if notes_text:
                    text_content_for_slide += f"\n\n--- Speaker Notes ---\n{notes_text}"
                
                contents_for_api.append(text_content_for_slide)

            doc_for_images.close()
        
        else:
            # Handle all other file types (PDF, DOCX, TXT) with text extraction
            app.logger.info(f"Processing text-based file: {file.filename}")
            extracted_text = ""
            file_stream = io.BytesIO(file_bytes)

            if filename_lower.endswith('.pdf'):
                with fitz.open(stream=file_stream, filetype="pdf") as doc:
                    for page in doc: extracted_text += page.get_text() + "\n"
            elif filename_lower.endswith('.docx'):
                doc = Document(file_stream)
                for para in doc.paragraphs: extracted_text += para.text + "\n"
            elif filename_lower.endswith('.txt'):
                extracted_text = file_bytes.decode('utf-8', errors='ignore')
            else:
                return jsonify({"success": False, "message": "Unsupported file type. Please use PDF, TXT, DOCX, or PPTX."}), 400
            
            if not extracted_text.strip(): return jsonify({"success": False, "message": "Could not extract any text from the document."}), 400
            
            # For text-only files, the prompt structure is simpler
            user_prompt = f"Please generate {num_questions} questions from the following document.\n- Subject: {subject}\n- Grade Level: {grade}\n- Desired Question Types: {question_types}\n- Additional Teacher Notes: {notes}\n\nSource Text:\n---\n{extracted_text[:30000]}"
            contents_for_api = [SYSTEM_PROMPT, user_prompt]

    except Exception as e:
        app.logger.error(f"Error parsing form or file: {e}", exc_info=True)
        return jsonify({"success": False, "message": "Failed to process the uploaded document."}), 500

    # --- Step 5: AI Call ---
    try:
        client = genai.Client(api_key=api_key)
        generation_config = types.GenerateContentConfig(temperature=0.7, max_output_tokens=8192, response_mime_type="application/json")
        
        app.logger.info(f"Sending request to Gemini with {len(contents_for_api)} parts for user {user.id}.")
        response = client.models.generate_content(
            model='models/gemini-1.5-flash-latest',
            contents=contents_for_api,
            config=generation_config
        )
        generated_json = json.loads(response.text)
        
        if not is_dev_user:
            new_log = UsageLog(user_id=user.id)
            db.session.add(new_log)
            db.session.commit()
            app.logger.info(f"Logged successful generation for user {user.email}.")
        
        return jsonify({"success": True, "questions": generated_json})

    except Exception as e:
        app.logger.error(f"Error during AI call for user {user.id}: {e}", exc_info=True)
        # Check for specific Google API errors if needed
        if isinstance(e, google_exceptions.GoogleAPICallError):
            return jsonify(success=False, message=f"A Google API error occurred: {e.reason}"), 502 # Bad Gateway
        return jsonify(success=False, message=f"An error occurred with the AI service: {str(e)}"), 500

# --- Server Execution ---
if __name__ == '__main__':
    with app.app_context():
        app.logger.info("Checking for and creating database if it doesn't exist...")
        db.create_all()
        app.logger.info("Database is ready.")
    
    app.logger.info("Starting production server with Waitress...")
    serve(app, host='0.0.0.0', port=5000)